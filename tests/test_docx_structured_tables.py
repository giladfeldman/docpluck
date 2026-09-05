"""DOCX table extraction: `extract_docx_structured` (v2.4.139).

WHAT THIS CLOSES. `extract_pdf_structured` is PDF-only, so `tables[]` and
`flattened_rows[]` were empty for every DOCX and a consumer that verifies
statistics from table rows received nothing for that whole input format --
while a DOCX *states* its table grid exactly, in `w:tbl`. The paper's numbers
were sitting in the file, fully structured, and docpluck was throwing them away.

ENGINE CHOICE IS MEASURED, NOT ASSUMED: `docs/BENCHMARKS_docx_engines_2026-09.md`
scores five candidates over 16 English DOCX (5,091 truth cells, 3,908 of them
statistic-bearing) with a two-sided metric. mammoth wins at statistic-cell recall
1.0000 / 0 fabrications, and is the only candidate that adds no dependency.

These tests were written against a tree with no `extract_docx_structured` at all
and watched RED first (they failed at import). A test authored after the fix that
merely re-asserts current behaviour proves nothing.
"""

from __future__ import annotations

import io
import json
import os
import re
import zipfile
from pathlib import Path

import pytest

pytest.importorskip("mammoth", reason="mammoth not installed (pip install docpluck[docx])")
pytest.importorskip("bs4", reason="beautifulsoup4 not installed (pip install docpluck[html])")
pytest.importorskip("docx", reason="python-docx not installed (dev dependency)")

from docx import Document  # noqa: E402
from docx.shared import Pt  # noqa: E402

from docpluck import (  # noqa: E402
    extract_docx_structured,
    extract_pdf_structured,
    flatten_tables_for_paper,
)
from docpluck.tables.flatten import (  # noqa: E402
    _cells_to_grid,
    _clean_grid,
    flatten_table,
)


# ── Fixture builders ────────────────────────────────────────────────────────
#
# Synthetic DOCX here, real DOCX in the custody-corpus tests at the bottom.
# A synthetic fixture can only confirm the model it was built from, so it pins
# BEHAVIOUR (does the value reach the consumer?) while the real corpus pins
# PREVALENCE (does this shape occur, and how often?).


def _docx_with_stats_table() -> bytes:
    """A statistics table using the header vocabulary real DOCX actually use.

    THE HEADERS HERE ARE NOT INVENTED. An earlier draft of this fixture headed
    the test-statistic column `t(48)`, which fails to classify -- and the
    tempting fix was to add a `t_with_df` role to `tables/flatten.py` beside the
    existing `F_with_df`. Measured first, over every real DOCX available (55
    documents with table cells): **`t(df)` occurs 0 times, `chi2(df)` 0 times,
    `F(a,b)` once.** The real vocabulary is bare and already classified -- `p`
    (118 cells), `df` (38), `F` (35), `N` (30), `SD` (24), `SE` (23), `t` (17),
    `M` (13), `d` (13), `95% CI` (16).

    So the defect was the FIXTURE, not the library, and adding the rule would
    have been pure false-positive surface justified by a string this file made
    up -- the exact shape of the retired A3d rule (CLAUDE.md: case studies come
    from real papers; measure the denominator separately from the shape).

    Regenerate the counts with `tools/diag/docx_tool_benchmark.py`'s corpus and
    a header census before touching `_ROLE_PATTERNS` on this evidence.
    """
    d = Document()
    d.add_paragraph("Results")
    d.add_paragraph("Table 1. Descriptive statistics and test results by condition.")
    t = d.add_table(rows=3, cols=6)
    hdr = ["Condition", "M", "SD", "t", "df", "p"]
    for c, text in enumerate(hdr):
        t.cell(0, c).text = text
    for r, row in enumerate(
        [["Treatment", "4.21", "1.12", "3.45", "48", ".001"],
         ["Control", "3.42", "0.89", "1.02", "48", ".312"]],
        start=1,
    ):
        for c, text in enumerate(row):
            t.cell(r, c).text = text
    d.add_paragraph("Note. N = 50 per condition. Cohen's d = 0.79.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _docx_with_merged_header() -> bytes:
    """A two-row header where the first column spans both rows.

    This is the shape that broke the benchmark's own adapter: ignore `rowspan`
    and the second header row slides left, so every value lands under the wrong
    statistic's name. `tables/flatten.py` binds values to headers by position,
    so a shift here is not cosmetic -- it publishes a t-statistic labelled `p`.
    """
    d = Document()
    d.add_paragraph("Table 2. Effect sizes by group.")
    t = d.add_table(rows=3, cols=3)
    t.cell(0, 0).merge(t.cell(1, 0))
    t.cell(0, 0).text = "Group"
    t.cell(0, 1).merge(t.cell(0, 2))
    t.cell(0, 1).text = "Effect"
    t.cell(1, 1).text = "d"
    t.cell(1, 2).text = "95% CI"
    t.cell(2, 0).text = "Younger"
    t.cell(2, 1).text = "0.42"
    t.cell(2, 2).text = "[0.18, 0.66]"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _docx_with_cell(text: str) -> bytes:
    d = Document()
    d.add_paragraph("Table 1. One value.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Estimate"
    t.cell(0, 1).text = "CI"
    t.cell(1, 0).text = "beta"
    t.cell(1, 1).text = text
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _docx_with_stacked_cell() -> bytes:
    """A cell holding two values as separate paragraphs.

    The benchmark's ground-truth reader fused exactly this shape into `.556.390`
    -- a number nobody printed. The library must not repeat it.
    """
    d = Document()
    d.add_paragraph("Table 1. Stacked.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Variable"
    t.cell(0, 1).text = "Loadings"
    t.cell(1, 0).text = "Item 1"
    cell = t.cell(1, 1)
    cell.text = ".556"
    cell.add_paragraph(".390")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _with_tracked_deletion(docx_bytes: bytes) -> bytes:
    """Rewrite one run into a `w:del` deletion, keeping the document valid."""
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as z:
        parts = {n: z.read(n) for n in z.namelist()}
        infos = {i.filename: i for i in z.infolist()}
    xml = parts["word/document.xml"].decode("utf-8")
    xml = xml.replace(
        "<w:t>3.45</w:t>",
        "<w:t>3.45</w:t></w:r>"
        '<w:del w:id="9001" w:author="a" w:date="2026-09-04T00:00:00Z">'
        '<w:r><w:delText>9.99</w:delText></w:r></w:del><w:r>',
        1,
    )
    parts["word/document.xml"] = xml.encode("utf-8")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        for name in parts:
            zo.writestr(infos[name], parts[name])
    return out.getvalue()


# ── The contract shape: additive, never renamed ─────────────────────────────


class TestResultShapeIsAdditive:
    """The DOCX result must be the SAME StructuredResult the PDF path returns.

    A consumer reading `result["tables"]` / `flattened_rows` must not need a
    second code path, and no existing field may be renamed. A rename that
    silently empties every downstream consumer is the most expensive defect this
    library has shipped, and it turned no test red anywhere.
    """

    def test_docx_result_has_exactly_the_pdf_result_keys(self):
        docx_keys = set(extract_docx_structured(_docx_with_stats_table()))
        pdf_keys = set(extract_pdf_structured(b"%PDF-1.4 not really a pdf"))
        assert docx_keys == pdf_keys, (
            "the DOCX StructuredResult drifted from the PDF one; a consumer "
            f"would need two code paths. only-docx={docx_keys - pdf_keys} "
            f"only-pdf={pdf_keys - docx_keys}"
        )

    def test_table_carries_every_declared_table_field(self):
        from docpluck.tables import Table

        res = extract_docx_structured(_docx_with_stats_table())
        assert res["tables"], "no table extracted from a document with one table"
        declared = set(Table.__annotations__)
        got = set(res["tables"][0])
        assert declared <= got, f"Table fields missing from the DOCX path: {declared - got}"

    def test_method_names_the_engine_that_actually_ran(self):
        res = extract_docx_structured(_docx_with_stats_table())
        assert "mammoth" in res["method"], (
            f"method={res['method']!r} does not name the engine; an unlabelled "
            "engine is the 'no pretending' violation"
        )


# ── The point of the whole exercise: statistics reach the consumer ──────────


class TestStatisticsReachTheConsumer:
    def test_cells_carry_the_published_values(self):
        res = extract_docx_structured(_docx_with_stats_table())
        texts = {c["text"] for t in res["tables"] for c in t["cells"]}
        for value in ("4.21", "1.12", "3.45", ".001", "3.42", "0.89"):
            assert value in texts, f"{value!r} was printed in the table and is not in cells"

    def test_flatten_binds_each_value_to_its_column_header(self):
        """`flatten_tables_for_paper` is the contract downstream verifiers read."""
        res = extract_docx_structured(_docx_with_stats_table())
        rows = flatten_tables_for_paper(res["tables"])
        assert rows, "a 5-column statistics table produced no FlattenedRow"
        treatment = [r for r in rows if r["row_label"] == "Treatment"]
        assert treatment, f"no row labelled 'Treatment'; got {[r['row_label'] for r in rows]}"
        fields = treatment[0]["fields"]
        assert fields.get("M") == 4.21, f"M not bound: {fields}"
        assert fields.get("SD") == 1.12, f"SD not bound: {fields}"
        assert fields.get("t") == 3.45, f"t not bound: {fields}"
        assert fields.get("df") == 48, f"df not bound: {fields}"
        assert fields.get("p") == 0.001, f"p not bound: {fields}"

    def test_caption_and_label_are_recovered_from_the_paragraph_above(self):
        res = extract_docx_structured(_docx_with_stats_table())
        t = res["tables"][0]
        assert t["label"] == "Table 1", f"label={t['label']!r}"
        assert t["caption"] and "Descriptive statistics" in t["caption"], (
            f"caption={t['caption']!r} -- flatten reads caption vocabulary to type "
            "an unlabelled effect column, so losing it loses effect-size typing"
        )

    def test_html_is_emitted_for_a_structured_table(self):
        res = extract_docx_structured(_docx_with_stats_table())
        html = res["tables"][0]["html"]
        assert html and "<table>" in html


class TestCaptionRecovery:
    """Captions are not decoration: `flatten_table` reads caption + footnote
    vocabulary to type an unlabelled estimate column. A missing caption loses
    effect-size typing; a WRONG one types a statistic the table never reported,
    so the negative controls below matter more than the recall.

    Measured over the 87 tables of 13 real DOCX: matching only the PDF path's
    `Table \\d+` shape recovered **16.1%**. Supplementary labels (`Table S1`) and
    APA 7's two-paragraph label/title form account for nearly all the rest;
    handling both, with the double-assignment fixed, gives **67.8%**.
    """

    def _docx_caption(self, before: list[str], after: list[str]) -> bytes:
        d = Document()
        for p in before:
            d.add_paragraph(p)
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "Group"
        t.cell(0, 1).text = "d"
        t.cell(1, 0).text = "A"
        t.cell(1, 1).text = "0.42"
        for p in after:
            d.add_paragraph(p)
        buf = io.BytesIO()
        d.save(buf)
        return buf.getvalue()

    def test_supplementary_label_is_recovered(self):
        """`Table S1` is not `Table \\d+`, and DOCX is the format supplementary
        material is submitted in."""
        res = extract_docx_structured(
            self._docx_caption(["Table S1. Baseline characteristics."], [])
        )
        assert res["tables"][0]["label"] == "Table S1"

    def test_apa_two_paragraph_label_and_title_are_joined(self):
        res = extract_docx_structured(
            self._docx_caption(["Table 3", "Effect sizes by condition"], [])
        )
        t = res["tables"][0]
        assert t["label"] == "Table 3"
        assert "Effect sizes by condition" in (t["caption"] or ""), f"{t['caption']!r}"

    def test_a_caption_below_the_table_is_recovered(self):
        res = extract_docx_structured(
            self._docx_caption(["Some prose about the study."],
                               ["Table 4. Results by arm."])
        )
        assert res["tables"][0]["label"] == "Table 4"

    def test_a_paragraph_that_merely_refers_to_a_table_is_not_its_caption(self):
        """NEGATIVE CONTROL. Proximity must never promote body prose."""
        res = extract_docx_structured(
            self._docx_caption(["The effect was robust, as shown in Table 2 below."], [])
        )
        t = res["tables"][0]
        assert t["label"] is None and t["caption"] is None, (
            f"a body reference became a caption: {t['caption']!r}. It would then "
            "be fed to flatten's effect-type vocabulary."
        )

    def test_one_caption_paragraph_is_not_claimed_by_two_tables(self):
        """NEGATIVE CONTROL. A caption between two tables was claimed by the
        table above it (as its 'next') AND the one below it (as its 'previous'),
        so 3 real documents emitted one label on two grids -- a consumer joining
        on `label` would merge two unrelated tables."""
        d = Document()
        for i, label in enumerate(("Table 1. First.", "Table 2. Second."), start=1):
            d.add_paragraph(label)
            t = d.add_table(rows=2, cols=2)
            t.cell(0, 0).text = "Group"
            t.cell(0, 1).text = "d"
            t.cell(1, 0).text = f"A{i}"
            t.cell(1, 1).text = f"0.4{i}"
        buf = io.BytesIO()
        d.save(buf)
        labels = [t["label"] for t in extract_docx_structured(buf.getvalue())["tables"]]
        assert len(labels) == len(set(labels)), f"a label was assigned twice: {labels}"
        assert labels == ["Table 1", "Table 2"], labels


class TestTheHtmlChannelAgreesWithTheCells:
    """One input, one answer — across every field the same Table exposes.

    Found 2026-09-05 by the Sol and Grok seats of the release consult round;
    the Sonnet seat filed an explicit all-clear on this exact question, which
    was a WRONG REJECT (it reasoned in plan mode and never ran the code).

    `docx_tables` cleaned each cell with the DOCX chain at construction and then
    handed the grid to `cells_to_html`, whose `_html_escape` re-cleaned it with
    the **PDF** chain. So one DOCX cell produced two different values:

        cells[].text / raw_text   [20.45, 20.06]      as the author typed it
        html                      [-0.45, -0.06]      two minus signs invented

    The module's own docstring named `[20.45, 20.06]` as the misfire it
    prevented. It prevented it in the cells and not in the HTML — which is the
    "no pretending" failure and the "one concept, one table" failure at once.

    These tests assert the INVARIANT rather than the one string, because the
    string was already covered and the invariant is what was actually broken.
    """

    def _html_cells(self, docx_bytes: bytes) -> list[str]:
        html = extract_docx_structured(docx_bytes)["tables"][0]["html"]
        return re.findall(r"<t[dh]>(.*?)</t[dh]>", html, re.S)

    @pytest.mark.parametrize("value", [
        "[20.45, 20.06]",   # a PDF '2'-for-U+2212 shape that a DOCX author really typed
        r"\.001",           # a PDF '<'-as-backslash shape
        "Direction 3 manipulated attribute",  # a PDF 'x'-as-'3' shape
    ])
    def test_html_reports_the_same_text_the_cells_do(self, value):
        docx = _docx_with_cell(value)
        cells = {c["text"] for t in extract_docx_structured(docx)["tables"]
                 for c in t["cells"]}
        assert value in cells, f"the cell channel already lost it: {cells}"
        html_cells = {h.strip() for h in self._html_cells(docx)}
        assert value in html_cells, (
            f"html says {html_cells} but cells say {value!r}. One input, two "
            "answers — a consumer reading `html` and one reading `cells` "
            "disagree about what the paper printed."
        )

    def test_the_inferential_ci_repair_does_not_run_on_a_docx(self):
        """`-.73` beside `[-0.78, 0.67]` must NOT become `[-0.78, -0.67]`.

        That repair argues the interval must be negative because otherwise it
        excludes the estimate. Inferential evidence is the consumer's to act on
        (it holds the parsed statistic and has a UI); docpluck holds text and
        has no channel to announce that it guessed. Justified on a PDF, where a
        font really did drop the glyph — never on a DOCX, which carries real
        Unicode. Raised by the Grok seat, 2026-09-05.
        """
        d = Document()
        d.add_paragraph("Table 1. Correlations.")
        t = d.add_table(rows=2, cols=3)
        for c, text in enumerate(["Variable", "r", "95% CI"]):
            t.cell(0, c).text = text
        for c, text in enumerate(["Item 1", "-.73", "[-0.78, 0.67]"]):
            t.cell(1, c).text = text
        buf = io.BytesIO()
        d.save(buf)
        html = " ".join(self._html_cells(buf.getvalue()))
        assert "[-0.78, 0.67]" in html, f"the interval was rewritten: {html}"
        assert "-0.67" not in html, (
            "a minus sign was manufactured on the CI upper bound; the document "
            "does not contain it"
        )


class TestCaptionsAreNotStolenFromProse:
    """A wrong caption is worse than no caption.

    `flatten_table` reads caption vocabulary (`_effect_type_for`) to type an
    unlabelled estimate column, so one stray word publishes a real number under
    the wrong statistic's name. All three consult seats reproduced this
    independently on 2026-09-05 — the strongest agreement in the round.
    """

    def _docx(self, before: list[str]) -> bytes:
        d = Document()
        for p in before:
            d.add_paragraph(p)
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "Group"
        t.cell(0, 1).text = "Estimate"
        t.cell(1, 0).text = "Younger"
        t.cell(1, 1).text = "0.42"
        buf = io.BytesIO()
        d.save(buf)
        return buf.getvalue()

    def test_an_apa_title_that_names_another_table_is_not_this_ones_title(self):
        """The reproduced case: a `Table 4` label, then prose naming Table 3."""
        res = extract_docx_structured(self._docx(
            ["Table 4",
             "Table 3 reports the partial eta-squared for each comparison."]
        ))
        t = res["tables"][0]
        assert t["label"] == "Table 4"
        assert "eta" not in (t["caption"] or "").lower(), (
            f"prose naming another table became this one's title: {t['caption']!r}"
        )
        rows = flatten_tables_for_paper(res["tables"])
        assert "eta2" not in rows[0]["fields"], (
            f"a Cohen's d estimate was typed as eta-squared: {rows[0]['fields']}"
        )

    @pytest.mark.parametrize("prose", [
        "cf. Table 2 for the full set of comparisons.",
        "See Tables 2 and 3 for the partial eta-squared values.",
        "The effect held across conditions (Table 2).",
    ])
    def test_prose_that_mentions_a_table_never_becomes_one(self, prose):
        """These are the shapes the Grok seat raised — and the honest note is
        that `_TABLE_REFERENCE_RE` is NOT what rejects them.

        Mutation-checked 2026-09-05: with `_TABLE_REFERENCE_RE` reverted IN
        MEMORY to its original narrow form (a short verb list plus a singular
        `Table`), all three of these STILL return `label=None`. Two other
        things reject them first — `_DOCX_TABLE_LABEL_RE` is `$`-anchored, so
        a sentence is not a label paragraph, and `_NAMES_A_TABLE_RE` refuses
        any intervening paragraph that names a table as a title. The widened
        reference regex is belt-and-braces, **measured redundant on these
        inputs**.

        Recorded rather than quietly kept, because a guard nobody can show
        biting is a claim, and a test that passes for a reason other than the
        one it names is the wrong-reject class this project keeps finding.
        What this test pins is the OUTCOME — prose never becomes a caption —
        which is the property worth protecting whichever guard delivers it.
        """
        t = extract_docx_structured(self._docx([prose]))["tables"][0]
        assert t["label"] is None and t["caption"] is None, (
            f"{prose!r} became a caption: {t['caption']!r}"
        )

    def test_a_real_two_paragraph_apa_caption_still_works(self):
        """Two-sided control: the guards must not have killed the feature."""
        t = extract_docx_structured(self._docx(
            ["Table 5", "Effect sizes by condition"]
        ))["tables"][0]
        assert t["label"] == "Table 5"
        assert "Effect sizes by condition" in (t["caption"] or "")


class TestMergedHeadersDoNotShiftColumns:
    def test_two_row_header_with_a_row_span_keeps_its_columns(self):
        res = extract_docx_structured(_docx_with_merged_header())
        cells = res["tables"][0]["cells"]
        grid: dict[tuple[int, int], str] = {(c["r"], c["c"]): c["text"] for c in cells}
        assert grid.get((1, 1)) == "d", (
            f"the sub-header slid: row1 = {[grid.get((1, i)) for i in range(3)]}. "
            "flatten binds by position, so this labels an effect size with the "
            "wrong statistic's name"
        )
        assert grid.get((1, 2)) == "95% CI"

    def test_the_value_row_still_lines_up_under_its_header(self):
        res = extract_docx_structured(_docx_with_merged_header())
        rows = flatten_tables_for_paper(res["tables"])
        assert rows, "merged-header table produced no rows"
        assert rows[0]["fields"].get("d") == 0.42, (
            f"d not bound under a merged header: {rows[0]['fields']}"
        )


class TestCellTextIsNotFusedOrInvented:
    def test_two_paragraphs_in_one_cell_do_not_fuse_into_a_new_number(self):
        res = extract_docx_structured(_docx_with_stacked_cell())
        texts = {c["text"] for t in res["tables"] for c in t["cells"]}
        assert ".556.390" not in texts, (
            "two stacked values fused into a number the paper never printed -- "
            "the same class `_linearize_omml` exists to prevent"
        )
        joined = " ".join(texts)
        assert ".556" in joined and ".390" in joined

    def test_unicode_minus_is_folded_to_ascii(self):
        """CLAUDE.md hard rule 4 / LESSONS L-004. Measured: 78 occurrences of
        U+2212 in the table cells of 2 of 13 real DOCX in the corpus."""
        res = extract_docx_structured(_docx_with_cell("−0.42"))
        texts = {c["text"] for t in res["tables"] for c in t["cells"]}
        assert "-0.42" in texts, f"U+2212 not folded: {texts}"

    def test_a_pdf_glyph_repair_never_fires_on_a_docx(self):
        """docpluck EXTRACTS, it does not FIX THE PAPER.

        `[20.45, 20.06]` in a PDF is a corrupted `[-0.45, -0.06]`, because a
        broken font mapped U+2212 to `2`. A DOCX carries real Unicode: there is
        no such font layer, so the same string is a value the authors actually
        typed and it must pass through verbatim. Measured over the corpus:
        **0 occurrences of `(cid:N)` or NUL in any DOCX table cell**, so the PDF
        recovery chain could only ever misfire here.
        """
        res = extract_docx_structured(_docx_with_cell("[20.45, 20.06]"))
        texts = {c["text"] for t in res["tables"] for c in t["cells"]}
        assert "[20.45, 20.06]" in texts, (
            f"a PDF glyph repair rewrote a DOCX cell: {texts}. That fabricates a "
            "minus sign the document does not contain."
        )
        assert "[-0.45, -0.06]" not in texts

    def test_deleted_tracked_change_text_is_not_resurrected(self):
        """Emitting `w:delText` is not recovery, it is resurrecting a retraction."""
        res = extract_docx_structured(_with_tracked_deletion(_docx_with_stats_table()))
        texts = {c["text"] for t in res["tables"] for c in t["cells"]}
        assert not any("9.99" in t for t in texts), (
            f"author-deleted text reached the consumer as live data: {texts}"
        )
        assert any("3.45" in t for t in texts), (
            "the surviving value was dropped along with the deletion"
        )


class TestFieldsWithNoDocxMeaningAreHonest:
    """A DOCX has no page model and no page geometry. Say so; never invent one."""

    def test_geometry_absence_is_named_not_silent(self):
        t = extract_docx_structured(_docx_with_stats_table())["tables"][0]
        assert t["cell_geometry"] and "docx" in t["cell_geometry"], (
            f"cell_geometry={t['cell_geometry']!r} -- a zero bbox with no stated "
            "reason is indistinguishable from real geometry that happens to be zero"
        )
        assert all(c["bbox"] == (0.0, 0.0, 0.0, 0.0) for c in t["cells"])

    def test_capture_quality_is_none_because_nothing_was_captured(self):
        t = extract_docx_structured(_docx_with_stats_table())["tables"][0]
        for field in ("confidence", "accuracy", "whitespace", "camelot_flavor"):
            assert t[field] is None, (
                f"{field}={t[field]!r}: a DOCX table is STATED, not inferred, so "
                "there is no capture-quality signal to report. Emitting a number "
                "here would invent a confidence."
            )

    def test_page_is_zero_because_docx_has_no_pages(self):
        t = extract_docx_structured(_docx_with_stats_table())["tables"][0]
        assert t["page"] == 0, (
            f"page={t['page']!r}: OOXML has no page model -- pagination is decided "
            "by the renderer -- so any positive page number is fabricated."
        )

    def test_rendering_names_markup_rather_than_a_capture_heuristic(self):
        t = extract_docx_structured(_docx_with_stats_table())["tables"][0]
        assert t["rendering"] == "markup", (
            f"rendering={t['rendering']!r} claims a PDF capture path ran. "
            "Reusing 'lattice' here would be an unlabelled engine substitution."
        )


class TestDegradesWithoutLying:
    def test_a_document_with_no_tables_returns_an_empty_list_not_an_error(self):
        d = Document()
        d.add_paragraph("No tables here, just prose with t(10) = 2.1, p = .04.")
        buf = io.BytesIO()
        d.save(buf)
        res = extract_docx_structured(buf.getvalue())
        assert res["tables"] == []
        assert "t(10) = 2.1" in res["text"]

    def test_a_malformed_docx_reports_the_failure_rather_than_an_empty_success(self):
        with pytest.raises(Exception):
            extract_docx_structured(b"not a zip at all")


# ── Real documents: the synthetic fixtures cannot establish prevalence ──────


def _custody_docx() -> list[tuple[str, Path]]:
    repo = Path(
        os.environ.get("ARTICLE_REPOSITORY")
        or (Path(os.environ.get("VIBE_ROOT") or (Path.home() / "Vibe")) / "ArticleRepository")
    )
    index = repo / "index.json"
    if not index.exists():
        return []
    try:
        data = json.loads(index.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return []
    out = []
    for key, entry in data.items():
        if not isinstance(entry, dict):
            continue
        fn = str(entry.get("filename") or "")
        if not fn.lower().endswith(".docx"):
            continue
        p = repo / "fulltext" / fn
        if p.exists():
            out.append((key, p))
    return sorted(out)


@pytest.mark.skipif(not _custody_docx(), reason="no DOCX resolvable through article-finder")
def test_real_docx_statistic_cells_survive_to_the_flattened_rows():
    """Every statistic-bearing cell of a real paper must reach a consumer.

    Scored against the OOXML itself -- the DOCX states its own grid, so this is
    the primary source and not a second run of the extractor under test.
    """
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    import xml.etree.ElementTree as ET

    checked = 0
    for key, path in _custody_docx():
        b = path.read_bytes()
        with zipfile.ZipFile(io.BytesIO(b)) as z:
            if "word/document.xml" not in z.namelist():
                continue
            root = ET.fromstring(z.read("word/document.xml"))
        # Values the document states inside a table, read independently.
        stated: set[str] = set()
        for tbl in root.iter(f"{W}tbl"):
            if len(tbl.findall(f"{W}tr")) < 2:
                continue
            for tc in tbl.iter(f"{W}tc"):
                for p in tc.iter(f"{W}p"):
                    txt = "".join(
                        n.text or "" for n in p.iter(f"{W}t")
                    ).strip()
                    if re.fullmatch(r"[-−]?\d+\.\d+", txt):
                        stated.add(txt.replace("−", "-"))
        if len(stated) < 20:
            continue
        res = extract_docx_structured(b)
        got = {
            c["text"].replace("−", "-")
            for t in res["tables"]
            for c in t["cells"]
        }
        missing = stated - got
        assert len(missing) / len(stated) < 0.02, (
            f"{key}: {len(missing)}/{len(stated)} stated table numbers did not "
            f"reach cells; e.g. {sorted(missing)[:8]}"
        )
        checked += 1
        if checked >= 4:
            break
    assert checked, (
        "no custody DOCX had >=20 decimal table values -- this assertion ran "
        "against nothing, which is a false green, not a pass"
    )


# ── A DECLARED header count is a fact the file states, not a guess to override ──


def test_clean_grid_does_not_promote_past_a_declared_header_count():
    """`_clean_grid` must not out-vote a header count the DOCX itself declares.

    Word records repeating header rows in `w:trPr/w:tblHeader`; mammoth turns
    exactly that into `<th>` (mammoth/body_xml.py:372) and nothing else, so a
    non-zero `header_rows` on a `rendering="markup"` table is a STATED FACT.
    `_is_header_like_row` re-derives it from cell length and numeric ratio and
    silently promotes real data rows -- the deletion class this project ranks
    worst, because the rows do not appear anywhere in `flattened_rows`.

    Reproduced on a real document (see the custody test below); this unit pins
    the mechanism. The shape is 10.5281/zenodo.21911212 Table 10: an APA
    "M [95% CI]" row whose cells are values-with-intervals, which
    `_DATA_VALUE_CELL_RE` does not full-match, so the row reads as header text.
    """
    grid = [
        ["Measure", "Group", "Pre", "Post", "Change", "p"],
        ["Numeracy", "CG", "15.04 [13.93, 16.16]", "14.91 [13.33, 16.48]",
         "-0.13 [-2.05, 1.79]", ".891"],
        ["Numeracy", "PG", "14.32 [13.19, 15.45]", "22.09 [20.50, 23.67]",
         "7.77 [5.85, 9.69]", "< .001"],
        ["Numeracy", "EXG", "14.44 [13.31, 15.57]", "22.08 [20.49, 23.66]",
         "7.63 [5.71, 9.55]", "< .001"],
    ]
    # Without the declaration the heuristic eats the first two data rows.
    _hdr_guessed, body_guessed = _clean_grid(grid)
    assert len(body_guessed) == 1, (
        "control: this grid must reproduce the over-promotion, otherwise the "
        "assertion below passes for the wrong reason"
    )

    hdr, body = _clean_grid(grid, declared_header_rows=1)
    assert len(hdr) == 1, f"declared 1 header row, got {len(hdr)}"
    assert len(body) == 3, f"declared 1 header row, expected 3 body rows, got {len(body)}"
    assert any("7.77 [5.85, 9.69]" in c for r in body for c in r), (
        "the PG arm's change score and its interval must reach the body"
    )


def test_a_declared_header_count_is_a_ceiling_never_a_floor():
    """The declaration may only KEEP rows as data, never promote more to header.

    A ceiling can add body rows and can never remove one, so the change cannot
    introduce the very deletion it exists to prevent. A document declaring 3
    header rows on a grid the heuristic reads as 1 must still yield 1.
    """
    grid = [
        ["Outcome", "t", "p"],
        ["Acc", "2.10", ".04"],
        ["Speed", "1.30", ".30"],
    ]
    hdr_plain, body_plain = _clean_grid(grid)
    hdr_decl, body_decl = _clean_grid(grid, declared_header_rows=3)
    assert (len(hdr_decl), len(body_decl)) == (len(hdr_plain), len(body_plain)), (
        "a declaration LARGER than the heuristic must not promote extra rows"
    )
    assert len(body_decl) == 2


@pytest.mark.skipif(not _custody_docx(), reason="no DOCX resolvable through article-finder")
def test_real_docx_declared_header_rows_are_honoured_end_to_end():
    """Measured over every custody DOCX: no table may flatten fewer body rows
    than its own declared header count allows.

    Before the fix this failed on 21 of 122 real tables, deleting 48 data rows
    that carry published statistics -- e.g. 10.5281/zenodo.21911212 Table 10,
    whose `Numeracy PG` row states `7.77 [5.85, 9.69]` and `p < .001` and which
    reached no consumer at all.
    """
    from docpluck.tables.docx_tables import extract_tables_docx

    checked = violations = 0
    detail: list[str] = []
    for key, path in _custody_docx():
        try:
            tables, _method = extract_tables_docx(path.read_bytes())
        except Exception:
            continue
        for t in tables:
            declared = int(t.get("header_rows") or 0)
            # Only tables where the AUTHOR declared it: `header_rows` defaults
            # to 1 when nothing was marked, and a default is not evidence.
            if not any(c["is_header"] for c in t["cells"]):
                continue
            if declared < 1:
                continue
            checked += 1
            rows = flatten_table(t)
            grid = _cells_to_grid(t["cells"])
            hdr, _body = _clean_grid(grid, declared_header_rows=declared)
            if rows and len(hdr) > declared:
                violations += 1
                detail.append(f"{key} {t['id']}: declared={declared} used={len(hdr)}")
    assert checked >= 20, (
        f"only {checked} tables declared a header count -- this assertion ran "
        "against too little to mean anything, which is a false green"
    )
    assert not violations, (
        f"{violations}/{checked} tables over-promoted past their declared header "
        f"count: {detail[:6]}"
    )


# ── The benchmark harness is a shipped component; its ground truth must not be
#    blind to a construct the library deliberately handles ──────────────────


def test_benchmark_truth_grid_reads_omml_like_the_shipped_path_does():
    """`docx_tool_benchmark._tc_text` must see OMML, or it scores the wrong mammoth.

    The shipped DOCX table path runs `_inline_omml_runs` before mammoth, because
    `mammoth` has no model of `m:oMath` and silently drops it -- which deletes the
    NAME of a statistic (`etap2`, `chi2`, `rho`). The harness's ground truth read
    `w:t` only, so an equation cell was EMPTY in truth while production emitted
    `etap2 = .361`. The harness would therefore charge the shipped engine with
    FABRICATING that value, and would penalise any OMML-preserving candidate.

    Measured 2026-09-05: 0 of 6,605 table cells in the 19 custody DOCX carry OMML
    (149 OMML elements do exist, in 3 of those documents, outside tables), so the
    blind spot cannot have moved the recorded engine ranking. It is fixed here
    because the harness is re-usable and the next corpus may differ.
    """
    import xml.etree.ElementTree as ET
    import sys as _sys

    _sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "tools" / "diag"))
    from docx_tool_benchmark import _tc_text

    M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    xml = (
        '<w:tc xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
        "<w:p><m:oMath><m:r><m:t>ηp2 = .361</m:t></m:r></m:oMath></w:p></w:tc>"
    )
    tc = ET.fromstring(xml)
    assert any(n.tag in (f"{M}oMath", f"{M}oMathPara") for n in tc.iter()), (
        "control: the fixture must actually contain OMML"
    )
    got = _tc_text(tc).strip()
    assert ".361" in got, (
        f"truth grid lost the OMML cell entirely: {got!r}. The shipped path emits "
        "'ηp2 = .361' for this cell, so the harness would score it as fabricated."
    )

    # NEGATIVE CONTROL: an ordinary w:t cell must be unchanged by the fix, and a
    # w:delText cell must still be dropped (emitting it resurrects a retraction).
    plain = ET.fromstring(
        '<w:tc xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        "<w:p><w:r><w:t>.556</w:t></w:r></w:p><w:p><w:r><w:t>.390</w:t></w:r></w:p></w:tc>"
    )
    assert _tc_text(plain).strip() == ".556 .390", _tc_text(plain)
    deleted = ET.fromstring(
        '<w:tc xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        "<w:p><w:del><w:r><w:delText>9.99</w:delText></w:r></w:del>"
        "<w:r><w:t>1.23</w:t></w:r></w:p></w:tc>"
    )
    assert "9.99" not in _tc_text(deleted), _tc_text(deleted)
    assert "1.23" in _tc_text(deleted), _tc_text(deleted)
    assert W  # referenced for symmetry with the namespace pair above
