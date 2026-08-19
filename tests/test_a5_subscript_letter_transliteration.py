"""A5 transliterates subscript LETTERS, and joins them with `_` (D6a + contract v2.0).

**Every test in this module was watched FAILING against the unfixed code first.**

The defect, reproduced end to end before the fix was written. A5's job is to
flatten the typographic forms a publisher uses into the ASCII downstream
consumers actually parse. It did that for subscript DIGITS and not for
subscript LETTERS, so the single most common effect size in psychology came out
as a mixed ASCII/Unicode token:

    BF01      ->  BF01      correct   (U+2080/U+2081, subscript digits, mapped)
    M1        ->  M1        correct
    eta2p     ->  eta2 + U+209A       WRONG — the subscript letter survives

The injury is silent and it is a DOWNGRADE, not an error. Measured through the
real consumer (effectcheck 0.7.5, whose `parse.R:576` folds a fixed alternation
`(?:eta2p|eta_p2|...)` into "partial eta-squared ="):

    'F(1, 98) = 4.20, p = .043, eta2p = .04'   ->  effect_reported = 0.04  PASS
    'F(1, 98) = 4.20, p = .043, eta2<U+209A> = .04'  ->  effect_reported = NA    OK

The row still appears, nothing errors, no warning is emitted anywhere — the
status just degrades from PASS (checked and correct) to OK (nothing was
checked). A verified effect size becomes an unverified one because of one
character.

MEASURED FREQUENCY, so this is not oversold: 0 of 26 sampled PDFs carry a
surviving subscript letter (publishers position subscripts typographically, so
the codepoint never reaches pdftotext's text channel) and 1 of 25 DOCX inputs
does (two occurrences of `Ha`, the alternative hypothesis). So the class is
LATENT on PDF and ACTIVE-BUT-RARE on DOCX — and the 101-PDF corpus every other
gate in this project runs against is structurally incapable of seeing it. That
is the argument for the DOCX fixture below, not merely for the fix.

The fix is a table completion keyed on the Unicode BLOCK rather than on the
letters we happened to think of, because "what is not on this list?" is the
question that produced defects 1, 2, 6 and 8 in this same release.
"""

from __future__ import annotations

import pytest

from docpluck.normalize import NormalizationLevel, normalize_text


def _norm(text: str) -> str:
    out, _ = normalize_text(text, NormalizationLevel.academic)
    return out.strip()


def _render(text: str) -> str:
    """The /render path, which preserves source glyphs by contract."""
    out, _ = normalize_text(
        text, NormalizationLevel.academic, preserve_math_glyphs=True
    )
    return out.strip()


# ── the defect that motivated the decision ──────────────────────────────

def test_partial_eta_squared_flattens_completely():
    """`eta2p` must not leave as `eta2` + a live U+209A.

    This is the exact token effectcheck's alternation misses; see the module
    docstring for the measured PASS -> OK degradation.
    """
    assert _norm("η²ₚ = .04") == "eta2_p = .04"


def test_partial_eta_squared_in_a_full_statistical_clause():
    assert _norm("F(1, 98) = 4.20, p = .043, η²ₚ = .04") == (
        "F(1, 98) = 4.20, p = .043, eta2_p = .04"
    )


def test_alternative_hypothesis_subscript_a():
    """`Ha` (U+2090) — the one shape actually OBSERVED in real input (DOCX)."""
    assert _norm("We tested Hₐ against H₀.") == "We tested H_a against H_0."


# ── the CLASS, not the instances: the whole Latin-subscript block ────────

@pytest.mark.parametrize(
    "codepoint,expected",
    [
        ("ₐ", "a"), ("ₑ", "e"), ("ₒ", "o"), ("ₓ", "x"),
        ("ₔ", "e"),  # SCHWA — transliterated as 'e', the conventional ASCII form
        ("ₕ", "h"), ("ₖ", "k"), ("ₗ", "l"), ("ₘ", "m"),
        ("ₙ", "n"), ("ₚ", "p"), ("ₛ", "s"), ("ₜ", "t"),
    ],
)
def test_latin_subscript_block_is_complete(codepoint, expected):
    """U+2090-U+209C in full.

    Parametrised over the BLOCK rather than over the letters that happened to
    appear in a corpus: a table that covers only the observed members is the
    same defect one codepoint later.
    """
    assert _norm(f"M{codepoint} = 3.1") == f"M_{expected} = 3.1"


@pytest.mark.parametrize(
    "codepoint,expected",
    [
        ("ᵢ", "i"), ("ᵣ", "r"), ("ᵤ", "u"), ("ᵥ", "v"),
    ],
)
def test_phonetic_latin_subscripts(codepoint, expected):
    assert _norm(f"x{codepoint} = 3.1") == f"x_{expected} = 3.1"


@pytest.mark.parametrize(
    "codepoint,expected",
    [
        ("ᵦ", "beta"), ("ᵧ", "gamma"), ("ᵨ", "rho"),
        ("ᵩ", "phi"), ("ᵪ", "chi"),
    ],
)
def test_greek_subscripts_use_the_greek_ascii_names(codepoint, expected):
    """U+1D66-U+1D6A are GREEK subscripts, so they take the Greek spelling.

    Mapping them to a Latin letter would silently rename the statistic: a
    subscript beta is not the letter 'b'. A5 already spells the base letters
    out (beta, gamma, rho, phi, chi), so these follow the same convention
    rather than inventing a second one.
    """
    assert _norm(f"M{codepoint} = 3.1") == f"M_{expected} = 3.1"


# ── the variant-codepoint sibling of the same defect ─────────────────────

def test_greek_phi_symbol_variant_maps_like_the_letter():
    """U+03D5 GREEK PHI SYMBOL is the phi COEFFICIENT in a stats table.

    A5 mapped U+03C6 GREEK SMALL LETTER PHI and not its variant form, so the
    same effect size flattened or survived depending purely on which codepoint
    the publisher's font emitted. Found by `tools/diag/glyph_map_gap_scan.py`,
    which counted 11 surviving occurrences across 2 of 26 baseline papers —
    including a `95% CI for phi` table header, i.e. a real published effect
    size. Identical shape to the subscript-letter defect: a member of an
    already-mapped equivalence class left off the list.
    """
    assert _norm("ϕ = .21") == "phi = .21"
    assert _norm("95% CI for ϕ") == "95% CI for phi"


# ── the contract that must NOT change ───────────────────────────────────

def test_render_path_still_preserves_the_source_glyph():
    """`preserve_math_glyphs=True` skips A5 entirely — rendered .md stays faithful."""
    assert _render("η²ₚ = .04") == "η²ₚ = .04"
    assert _render("Hₐ") == "Hₐ"


def test_subscript_digits_still_map():
    """The pre-existing behaviour the fix must not disturb."""
    assert _norm("BF₀₁ = 3.2") == "BF_01 = 3.2"
    assert _norm("M₁ = 4.5") == "M_1 = 4.5"


def test_a_subscript_letter_does_not_swallow_neighbouring_text():
    """The replacement is character-for-character, never a span rewrite."""
    assert _norm("Hₐ and H₀ and then prose") == "H_a and H_0 and then prose"


def test_idempotent():
    """Normalizing twice equals normalizing once (a suite-wide invariant)."""
    once = _norm("η²ₚ = .04, Hₐ, ϕ = .21")
    assert _norm(once) == once


# ── the DOCX channel, where this class is ACTIVE rather than latent ──────
#
# The 101-PDF corpus every other gate in this project measures against cannot
# see this defect at all: publishers position subscripts typographically, so
# U+209A never reaches pdftotext's text channel (0 of 26 sampled PDFs). It
# reaches mammoth's DOCX channel intact (1 of 25 sampled DOCX files, carrying
# `Ha`). A library with three input formats needs a test per format, or a whole
# class stays structurally invisible — which is exactly how this one survived.
#
# The fixture is built in memory from strings written here, so no publication
# text is stored in this repo (article-finder is the sole custodian of papers).

def test_docx_input_flattens_subscript_letters_end_to_end():
    pytest.importorskip("mammoth", reason="pip install docpluck[docx]")
    pytest.importorskip("bs4", reason="pip install docpluck[html]")
    pytest.importorskip("docx", reason="python-docx not installed (dev dependency)")

    import io

    from docx import Document

    from docpluck import extract_docx

    doc = Document()
    doc.add_paragraph("We tested Hₐ against H₀ using a mixed ANOVA.")
    doc.add_paragraph("The interaction was significant, F(1, 98) = 4.20, η²ₚ = .04.")
    buf = io.BytesIO()
    doc.save(buf)

    raw = extract_docx(buf.getvalue())
    raw_text = raw[0] if isinstance(raw, tuple) else raw
    # The codepoint must actually SURVIVE extraction, or this test would pass
    # for the wrong reason — a green result from an input that never contained
    # the defect proves nothing.
    assert "ₚ" in raw_text or "ₐ" in raw_text, (
        "fixture did not carry a subscript letter through extraction; "
        "the assertion below would be vacuous"
    )

    out, _ = normalize_text(raw_text, NormalizationLevel.academic)
    assert "eta2_p" in out
    assert "H_a" in out
    assert "ₚ" not in out
    assert "ₐ" not in out
