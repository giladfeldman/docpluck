"""A crash in DOCX table extraction must NAME itself in the artifact.

Why this file exists (rule 40, v2.4.144 review, 2026-09-24).
`_extract_docx_structured` catches any exception from `extract_tables_docx` so a
working text extraction is never lost, and records it twice:
``tables_failed:<Exc>`` in ``method`` and the ``docx_table_extraction_exception``
fallback. Nothing asserted either. An empty ``tables`` with no token is
indistinguishable from a document that has no tables -- the false green this
project keeps finding -- so this gate fails if the recording ever stops.

A synthetic DOCX is enough here: this path reads no layout, and the claim under
test is the recording of an injected exception, not the prevalence of any shape
in real papers.
"""

import importlib
import io

import pytest

pytest.importorskip("mammoth", reason="mammoth not installed (pip install docpluck[docx])")
pytest.importorskip("docx", reason="python-docx not installed (dev dependency)")

from docx import Document

from docpluck import extract_docx_structured

# NOT `import docpluck.extract_docx_structured as eds`: the package re-exports the
# FUNCTION under the module's own name, so that form binds the function and the
# monkeypatch below would target the wrong object.
eds = importlib.import_module("docpluck.extract_docx_structured")


@pytest.fixture(scope="module")
def docx_bytes() -> bytes:
    d = Document()
    d.add_paragraph("Results")
    d.add_paragraph("Table 1. Descriptive statistics by condition.")
    t = d.add_table(rows=3, cols=4)
    for r, row in enumerate(
        [["Condition", "M", "SD", "p"],
         ["Treatment", "4.21", "1.12", ".001"],
         ["Control", "3.42", "0.89", ".312"]]
    ):
        for c, text in enumerate(row):
            t.cell(r, c).text = text
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_control_the_table_pass_runs_and_records_nothing(docx_bytes):
    result = extract_docx_structured(docx_bytes)
    method = result["method"]
    assert result["tables"], f"no table extracted: the path is not reached. method={method}"
    assert "+tables:" in method, method
    assert "tables_failed:" not in method, method
    assert "docx_table_extraction_exception" not in result["fallbacks"]


def test_a_crashing_table_pass_is_named_in_method_and_fallbacks(docx_bytes, monkeypatch):
    def _boom(*_a, **_kw):
        raise RuntimeError("injected: docx table extraction crashed")

    # The module binds the name at import, so patch it where it is looked up.
    monkeypatch.setattr(eds, "extract_tables_docx", _boom)
    result = extract_docx_structured(docx_bytes)
    method = result["method"]

    assert result["tables"] == []
    assert "Treatment" in result["text"], "the text extraction must survive a table crash"
    assert "tables_failed:RuntimeError" in method, (
        "the table pass crashed and the method string reads like a document with no "
        f"tables. method={method}"
    )
    assert "docx_table_extraction_exception" in result["fallbacks"], (
        f"table crash not recorded. fallbacks={result['fallbacks']}"
    )
